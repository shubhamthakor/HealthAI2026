import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ----------------------------------------------------
# PRESENTATION CONFIGURATION
# ----------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6] # Blank slide layout

# ----------------------------------------------------
# DESIGN CONSTANTS (Modern Healthcare Theme)
# ----------------------------------------------------
COLOR_SLATE_900 = RGBColor(15, 23, 42)    # Dark slate for primary titles
COLOR_SLATE_600 = RGBColor(71, 85, 105)   # Slate gray for normal text
COLOR_TEAL_600 = RGBColor(13, 148, 136)   # Teal for primary branding
COLOR_SKY_500 = RGBColor(14, 165, 233)    # Sky blue for accents
COLOR_SLATE_50 = RGBColor(248, 250, 252)  # Soft background gray for cards
COLOR_SLATE_200 = RGBColor(226, 232, 240) # Thin border gray
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_RED_500 = RGBColor(239, 68, 68)     # Accent red
COLOR_GREEN_500 = RGBColor(34, 197, 94)   # Accent green

# ----------------------------------------------------
# LAYOUT HELPERS
# ----------------------------------------------------
def apply_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE

def add_header(slide, title, category="PROJECT OVERVIEW"):
    # Category Tracker
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.name = "Calibri"
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_TEAL_600
    
    # Title Header
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.733), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_SLATE_900

    # Underline Separator
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.015))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_SLATE_200
    shape.line.fill.background()

def add_card(slide, left, top, width, height, title, bullet_points, accent_color=COLOR_TEAL_600):
    # Rounded Card Background
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_SLATE_50
    card.line.color.rgb = COLOR_SLATE_200
    card.line.width = Pt(1)

    # Accent Strip on Left Border
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    # Card Title
    title_box = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.15), width - Inches(0.25), Inches(0.4))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(14)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_SLATE_900

    # Card Content
    content_box = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.55), width - Inches(0.25), height - Inches(0.7))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    tf_content.margin_left = tf_content.margin_right = tf_content.margin_top = tf_content.margin_bottom = 0
    
    for i, pt in enumerate(bullet_points):
        p = tf_content.paragraphs[0] if i == 0 else tf_content.add_paragraph()
        p.text = f"•  {pt}"
        p.font.name = "Calibri"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_SLATE_600
        p.space_after = Pt(4)

# ----------------------------------------------------
# SLIDE 1: TITLE SLIDE (Project Introduction)
# ----------------------------------------------------
slide1 = prs.slides.add_slide(blank_layout)
apply_background(slide1)

# Left Column - Title Typography
left_col = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(5.5), Inches(5.0))
tf_left = left_col.text_frame
tf_left.word_wrap = True

# Category Track
p_track = tf_left.paragraphs[0]
p_track.text = "MAJOR PROJECT PRESENTATION"
p_track.font.name = "Calibri"
p_track.font.size = Pt(11)
p_track.font.bold = True
p_track.font.color.rgb = COLOR_TEAL_600
p_track.space_after = Pt(20)

# Project Name
p_name = tf_left.add_paragraph()
p_name.text = "HealthAI"
p_name.font.name = "Calibri"
p_name.font.size = Pt(54)
p_name.font.bold = True
p_name.font.color.rgb = COLOR_TEAL_600
p_name.space_after = Pt(6)

# Main Title Description
p_subtitle = tf_left.add_paragraph()
p_subtitle.text = "AI Based Disease Detection and Smart Appointment Booking System"
p_subtitle.font.name = "Calibri"
p_subtitle.font.size = Pt(20)
p_subtitle.font.bold = True
p_subtitle.font.color.rgb = COLOR_SLATE_900
p_subtitle.space_after = Pt(24)

# Tech Pill Indicators
p_tech = tf_left.add_paragraph()
p_tech.text = "React.js  |  Node.js (Express)  |  FastAPI  |  MongoDB Atlas"
p_tech.font.name = "Calibri"
p_tech.font.size = Pt(12)
p_tech.font.bold = True
p_tech.font.color.rgb = COLOR_SKY_500
p_tech.space_after = Pt(36)

# Presenter
p_pres = tf_left.add_paragraph()
p_pres.text = "DEVELOPED BY:\nShubham Thakor"
p_pres.font.name = "Calibri"
p_pres.font.size = Pt(13)
p_pres.font.bold = True
p_pres.font.color.rgb = COLOR_SLATE_900

# Right Column - Introduction Cards
# Card 1: What is HealthAI
add_card(slide1, Inches(6.8), Inches(1.0), Inches(5.7), Inches(1.8), 
         "What is HealthAI?", 
         ["AI-powered diagnostic system matching symptoms to diseases.",
          "Recommends targeted specialists based on classification models.",
          "Features real-time clinic queuing updates and booking dashboards."],
         COLOR_TEAL_600)

# Card 2: Problem Statement
add_card(slide1, Inches(6.8), Inches(3.0), Inches(5.7), Inches(1.8), 
         "The Healthcare Problem", 
         ["Patients routinely select incorrect medical specialists, delaying care.",
          "Lack of active estimated waiting times leads to crowded clinic rooms.",
          "Manual appointment scheduling processes lack transparent scheduling."],
         COLOR_RED_500)

# Card 3: Proposed Solution
add_card(slide1, Inches(6.8), Inches(5.0), Inches(5.7), Inches(1.8), 
         "Our Proposed Solution", 
         ["Groq Cloud NLP symptom translation across three language sets.",
          "Asynchronous FastAPI Random Forest classifier diagnostic evaluations.",
          "Socket.IO queue synchronization tracking checkup times dynamically."],
         COLOR_GREEN_500)


# ----------------------------------------------------
# SLIDE 2: PROJECT FEATURES
# ----------------------------------------------------
slide2 = prs.slides.add_slide(blank_layout)
apply_background(slide2)
add_header(slide2, "Comprehensive Project Features", "Core Platform Modules")

# Six features grid
features = [
    ("AI Disease Prediction", ["Scikit-Learn Random Forest Classifier inference.", "Categorical symptom vector evaluation.", "Translates user dialogue context via Groq API."], COLOR_TEAL_600),
    ("Doctor Discovery", ["Specialist matching from predicted illness keys.", "Availability filtering based on scheduled days.", "City-based query index maps across Gujarat."], COLOR_TEAL_600),
    ("Dynamic Queue Scheduler", ["Wait-time formula: Remaining patients * consultation time.", "Establishes shift start and end schedules.", "Prevents manual booking slot reservation blockages."], COLOR_SKY_500),
    ("Socket.IO Synchronization", ["Real-time event streams mapping patient dashboard counters.", "Updates waiting list updates on doctor completion calls.", "Synchronized checkup duration timer modifications."], COLOR_SKY_500),
    ("Shift Boundary Overflow Guard", ["Validates expected queue completion timeline dynamically.", "Automatically rejects bookings pushing past shift boundaries.", "Suggests next-day booking blocks to avoid queues."], COLOR_RED_500),
    ("Enterprise Notifications & Admin", ["Nodemailer email dispatch on booking confirmations.", "Leave manager updating clinic accessibility rules.", "Admin panel providing CRUD and query monitor logs."], COLOR_RED_500)
]

for idx, (title, pts, color) in enumerate(features):
    col = idx % 3
    row = idx // 3
    l = Inches(0.8) + col * Inches(3.9)
    t = Inches(1.8) + row * Inches(2.5)
    add_card(slide2, l, t, Inches(3.6), Inches(2.2), title, pts, color)


# ----------------------------------------------------
# SLIDE 3: TECHNOLOGY STACK
# ----------------------------------------------------
slide3 = prs.slides.add_slide(blank_layout)
apply_background(slide3)
add_header(slide3, "System Technology Stack Architecture", "Enterprise Technologies")

tech_stack = [
    ("Frontend UI / UX", ["React (Vite) - SPA rendering", "Tailwind CSS - Adaptive styles", "Axios - API client requests", "Socket.IO Client - Live feeds", "i18next - Interface translation"], COLOR_TEAL_600),
    ("Backend API Gateway", ["Node.js - High-throughput runtime", "Express.js - REST service endpoints", "JWT - Refresh/Access session tokens", "bcryptjs - Hash credential storage", "Helmet.js - Core headers security"], COLOR_TEAL_600),
    ("AI & Machine Learning", ["Python 3.10 - Scientific runtime", "FastAPI - Asynchronous microservice", "Scikit-Learn - Random Forest model", "Groq Llama 3 - NLP translations", "Pandas & NumPy - Preprocessing"], COLOR_SKY_500),
    ("Database & Operations", ["MongoDB Atlas - Cloud database service", "Mongoose - Document schema ODM", "Nodemailer / Resend - SMTP emails", "node-cron - Scheduled script runners", "Git / GitHub - Repository tracking"], COLOR_SKY_500)
]

for idx, (title, pts, color) in enumerate(tech_stack):
    l = Inches(0.8) + idx * Inches(2.93)
    add_card(slide3, l, Inches(1.8), Inches(2.73), Inches(4.8), title, pts, color)


# ----------------------------------------------------
# SLIDE 4: SYSTEM ARCHITECTURE
# ----------------------------------------------------
slide4 = prs.slides.add_slide(blank_layout)
apply_background(slide4)
add_header(slide4, "Multi-Service Architecture Topology", "System Integration")

# Draw visual block shapes for architecture
def draw_block_shape(slide, left, top, width, height, title, subtitle, color):
    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    block.fill.solid()
    block.fill.fore_color.rgb = COLOR_SLATE_50
    block.line.color.rgb = COLOR_SLATE_200
    block.line.width = Pt(1.5)
    
    # Left Border Accent
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.06), height)
    border.fill.solid()
    border.fill.fore_color.rgb = color
    border.line.fill.background()
    
    # Text Frame
    tb = slide.shapes.add_textbox(left + Inches(0.1), top + Inches(0.1), width - Inches(0.15), height - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.name = "Calibri"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_SLATE_900
    
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.name = "Calibri"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_SLATE_600
    p2.space_before = Pt(3)

# Column 1: Frontends
draw_block_shape(slide4, Inches(0.8), Inches(2.0), Inches(2.6), Inches(1.1), "Patient Portal", "React.js (Vite), i18n Translate, Speech Input", COLOR_TEAL_600)
draw_block_shape(slide4, Inches(0.8), Inches(3.4), Inches(2.6), Inches(1.1), "Doctor Console", "React.js (Vite), Socket.IO calculations", COLOR_TEAL_600)
draw_block_shape(slide4, Inches(0.8), Inches(4.8), Inches(2.6), Inches(1.1), "Admin Dashboard", "React.js (Vite), Doctor Manager, stats", COLOR_TEAL_600)

# Connectors Col 1 to Gateway
for top_inch in [2.45, 3.85, 5.25]:
    arrow = slide4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.55), Inches(top_inch), Inches(0.5), Inches(0.2))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = COLOR_SLATE_200
    arrow.line.fill.background()

# Column 2: Central API Gateway
draw_block_shape(slide4, Inches(4.2), Inches(2.9), Inches(3.2), Inches(2.1), "Central API Gateway", "Node.js & Express API Gateway\n• JWT Token validation & RBAC\n• Real-time waiting calculation engine\n• Nodemailer triggers", COLOR_SKY_500)

# Connectors Col 2 to Col 3
for top_inch in [2.35, 3.85, 5.35]:
    arrow = slide4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(7.55), Inches(top_inch), Inches(0.5), Inches(0.2))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = COLOR_SLATE_200
    arrow.line.fill.background()

# Column 3: Microservices & Data
draw_block_shape(slide4, Inches(8.2), Inches(1.8), Inches(4.3), Inches(1.2), "Database (MongoDB Atlas)", "Stores documents for Auth, user metadata, doctor profile shift rules, leaves, queue list booking states.", COLOR_TEAL_600)
draw_block_shape(slide4, Inches(8.2), Inches(3.3), Inches(4.3), Inches(1.2), "AI Services (Groq Llama 3 & Gemini)", "Intent extraction, conversational replies context formulation, multilingual translation.", COLOR_TEAL_600)
draw_block_shape(slide4, Inches(8.2), Inches(4.8), Inches(4.3), Inches(1.2), "FastAPI Microservice (Scikit-Learn)", "Python prediction API wrapping pre-trained Random Forest model running categorical feature prediction.", COLOR_SKY_500)


# ----------------------------------------------------
# SLIDE 5: BACKEND WORKFLOW
# ----------------------------------------------------
slide5 = prs.slides.add_slide(blank_layout)
apply_background(slide5)
add_header(slide5, "Central Backend Services Workflow", "System Workflow")

steps = [
    ("1. Handshake Auth", ["Public login & registration route handlers.", "Validates input sanitization metrics.", "Access and Refresh tokens setup.", "HttpOnly secure cookie configuration."], COLOR_TEAL_600),
    ("2. Route Guarding", ["Strict RBAC controller gates.", "Decodes JWT authentication payloads.", "Blocks out unauthorized role scopes.", "Applies DDoS request rate limiters."], COLOR_TEAL_600),
    ("3. Service Dispatch", ["Runs Groq symptom translator.", "Queries FastAPI ML disease predictions.", "Executes doctor location lookups.", "Validates dates constraint checks."], COLOR_SKY_500),
    ("4. Socket.IO Sync", ["Initializes client socket connections.", "Allocates room joins based on Doctor IDs.", "Pushes real-time wait recalculations.", "Triggers active consult room alerts."], COLOR_SKY_500),
    ("5. Notification Mailer", ["Fires transactional booking alerts.", "Configures custom HTML receipt layout.", "Runs node-cron leave resets.", "Logs API server errors to console."], COLOR_RED_500)
]

for idx, (title, pts, color) in enumerate(steps):
    l = Inches(0.8) + idx * Inches(2.36)
    add_card(slide5, l, Inches(2.0), Inches(2.16), Inches(4.5), title, pts, color)


# ----------------------------------------------------
# SLIDE 6: MACHINE LEARNING WORKFLOW
# ----------------------------------------------------
slide6 = prs.slides.add_slide(blank_layout)
apply_background(slide6)
add_header(slide6, "Machine Learning Diagnostic Pipeline", "Machine Learning Model")

# Left Column - Visual Step Pipeline
pipeline_steps = [
    ("1. Dataset Selection", "Parsed a database containing 132 binary categorical clinical symptoms."),
    ("2. Data Processing & Vectorization", "Processed sparse arrays using Pandas/NumPy, filtering duplicate inputs."),
    ("3. Model Ensemble Fitting", "Trained Scikit-Learn Random Forest Classifier with multiclass outcomes."),
    ("4. Serializing (Joblib)", "Compressed fitted weight matrices to binary payloads for instant server load."),
    ("5. FastAPI Inference Service", "Asynchronous endpoints hosting predictive models, responding in under 10ms.")
]

for idx, (step_title, step_desc) in enumerate(pipeline_steps):
    box_top = Inches(1.8) + idx * Inches(1.0)
    # Shape Rectangle
    r = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), box_top, Inches(5.4), Inches(0.85))
    r.fill.solid()
    r.fill.fore_color.rgb = COLOR_SLATE_50
    r.line.color.rgb = COLOR_SLATE_200
    r.line.width = Pt(1)
    
    # Accent indicator
    accent = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), box_top, Inches(0.06), Inches(0.85))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_TEAL_600
    accent.line.fill.background()
    
    # Text
    tb = slide6.shapes.add_textbox(Inches(1.0), box_top + Inches(0.08), Inches(5.0), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p1 = tf.paragraphs[0]
    p1.text = step_title
    p1.font.name = "Calibri"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_SLATE_900
    
    p2 = tf.add_paragraph()
    p2.text = step_desc
    p2.font.name = "Calibri"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_SLATE_600
    p2.space_before = Pt(2)

# Right Column - Why Random Forest
add_card(slide6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.85),
         "Why Scikit-Learn Random Forest Classifier?",
         ["Ensemble Architecture: Aggregates multiple decision trees to maintain robust generalization across sparse datasets.",
          "High Dimensional Generalization: Excelled at handling 132-dimension categorical feature arrays without overfitting.",
          "Deterministic Confidence Scores: Generates probability distributions (predict_proba) helping system filter diagnosis trust limits.",
          "Sub-10ms Server Latency: Ultra-light weight footprints make FastAPI inference highly compatible with interactive APIs."],
         COLOR_SKY_500)


# ----------------------------------------------------
# SLIDE 7: AI ASSISTANT WORKFLOW
# ----------------------------------------------------
slide7 = prs.slides.add_slide(blank_layout)
apply_background(slide7)
add_header(slide7, "Intelligent AI Agent Architecture", "AI Assistant Workflow")

# Left Column - Process Diagram
assistant_pipeline = [
    ("User Input Message", "Patient asks query: 'Book cardiologist appointment in Anand tomorrow'"),
    ("Intent Detection (Gemini API)", "Decodes query to match intent 'appointment_booking' and extracts variables."),
    ("Secure API Action Layer", "Node.js translates intent parameters and makes internal query checks."),
    ("Database Execution & Formatter", "Performs doctor lookups, books the queue, and writes response context.")
]

for idx, (step_title, step_desc) in enumerate(assistant_pipeline):
    box_top = Inches(1.8) + idx * Inches(1.2)
    r = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), box_top, Inches(5.4), Inches(0.95))
    r.fill.solid()
    r.fill.fore_color.rgb = COLOR_SLATE_50
    r.line.color.rgb = COLOR_SLATE_200
    r.line.width = Pt(1)
    
    accent = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), box_top, Inches(0.06), Inches(0.95))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_SKY_500
    accent.line.fill.background()
    
    tb = slide7.shapes.add_textbox(Inches(1.0), box_top + Inches(0.08), Inches(5.0), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p1 = tf.paragraphs[0]
    p1.text = step_title
    p1.font.name = "Calibri"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_SLATE_900
    
    p2 = tf.add_paragraph()
    p2.text = step_desc
    p2.font.name = "Calibri"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_SLATE_600
    p2.space_before = Pt(3)

# Right Column - NLP & Example
add_card(slide7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.2),
         "Multilingual Symptom Extraction (Groq)",
         ["Integrates Groq Llama-3-70B to parse conversational text from multilingual patients.",
          "Translates medical inputs in Hindi and Gujarati to clinical English symptom arrays.",
          "Guarantees language parity so localized input matches standard binary models."],
         COLOR_TEAL_600)

add_card(slide7, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.35),
         "Intent Parsing Example",
         ["User Input: 'મને તાવ છે (I have fever)'",
          "Intent Extracted: 'symptom_analysis'",
          "Extracted Entity: { symptoms: ['fever'], lang: 'gu' }",
          "Result: System translates symptom, runs Random Forest predict, maps matching General Physicians in the local city, and outputs card recommendations."],
         COLOR_TEAL_600)


# ----------------------------------------------------
# SLIDE 8: DATABASE DESIGN
# ----------------------------------------------------
slide8 = prs.slides.add_slide(blank_layout)
apply_background(slide8)
add_header(slide8, "MongoDB Atlas Schema Configurations", "Database Architecture")

# Left Column - Visual DB Collections Layout
db_cols = [
    ("Users Schema", "Unified collection tracking authentication. Models credentials, security details, and distinct role assignments ('patient' | 'doctor').", COLOR_TEAL_600),
    ("Doctors Schema", "Contains profile metrics: specializations, hospital location, city indexes, consultation fee rates, shift limits, and calendar availability.", COLOR_TEAL_600),
    ("Appointments Schema", "Tracks operational state: Mongoose ObjectId refs linking Patients and Doctors. Holds status history, wait calculations, and prescription details.", COLOR_SKY_500)
]

for idx, (title, desc, color) in enumerate(db_cols):
    box_top = Inches(1.8) + idx * Inches(1.6)
    
    r = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), box_top, Inches(5.4), Inches(1.4))
    r.fill.solid()
    r.fill.fore_color.rgb = COLOR_SLATE_50
    r.line.color.rgb = COLOR_SLATE_200
    r.line.width = Pt(1)
    
    accent = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), box_top, Inches(0.06), Inches(1.4))
    accent.fill.solid()
    accent.fill.fore_color.rgb = color
    accent.line.fill.background()
    
    tb = slide8.shapes.add_textbox(Inches(1.0), box_top + Inches(0.12), Inches(5.0), Inches(1.15))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.name = "Calibri"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_SLATE_900
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.name = "Calibri"
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_SLATE_600
    p2.space_before = Pt(4)

# Right Column - Data Operations
add_card(slide8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.6),
         "Operational Database Architecture",
         ["Schema Normalization: Relational links configured using Mongoose references rather than large nested embeds to preserve indexing speed.",
          "Indexing Strategy: Active database indices created on { email: 1 }, { specialization: 1 }, and { city: 1 } to keep discovery routes under 100ms.",
          "Prescription Models: Structured array fields mapping dosage intervals ('1-0-1') and instructions directly within the Appointment record.",
          "Doctor Leaves: Isolated collection storing doctor leaves to prevent scheduling conflicts during calendar searches."],
         COLOR_SKY_500)


# ----------------------------------------------------
# SLIDE 9: FUTURE SCOPE
# ----------------------------------------------------
slide9 = prs.slides.add_slide(blank_layout)
apply_background(slide9)
add_header(slide9, "Project Roadmap & Scalability Scope", "Future Scope")

future_scope = [
    ("Voice Assistant Integration", ["Direct speech-to-speech assistant.", "Allows scheduling dial-ins.", "Improves accessibility for users."], COLOR_TEAL_600),
    ("Medical Image Diagnostics", ["Integrating CNN models.", "Scans MRI/X-Ray image files.", "Identifies anomalies automatically."], COLOR_TEAL_600),
    ("OCR Prescription Scanner", ["Extracts physical handwriting text.", "Saves paper charts to MongoDB.", "Builds instant prescriptions."], COLOR_SKY_500),
    ("RAG Health Knowledge Base", ["LLM queries clinical books.", "Provides authenticated insights.", "Ensures safe clinical referencing."], COLOR_SKY_500),
    ("WebRTC Video Telemedicine", ["Virtual consulting dashboards.", "Prescription sharing during calls.", "Live browser consultation rooms."], COLOR_RED_500),
    ("Wearable Device IoT Sync", ["Pulls smartwatch vitals.", "Monitors heart rates automatically.", "Alerts doctors on abnormal vitals."], COLOR_RED_500)
]

for idx, (title, pts, color) in enumerate(future_scope):
    col = idx % 3
    row = idx // 3
    l = Inches(0.8) + col * Inches(3.9)
    t = Inches(1.8) + row * Inches(2.5)
    add_card(slide9, l, t, Inches(3.6), Inches(2.2), title, pts, color)


# ----------------------------------------------------
# SLIDE 10: CONCLUSION
# ----------------------------------------------------
slide10 = prs.slides.add_slide(blank_layout)
apply_background(slide10)
add_header(slide10, "Platform Summary & Future Outlook", "Project Conclusion")

# Left Column - Core Achievements Card
add_card(slide10, Inches(0.8), Inches(1.8), Inches(5.4), Inches(4.8),
         "Core Achievements",
         ["Successfully engineered a production-grade Web Portal integrating Machine Learning diagnostics and clinical NLP extraction.",
          "Developed an automated real-time Smart Queue Engine using wait-time formulas to resolve waiting line congestion.",
          "Hardened the multi-role gateway with secure Double-Token JWT rotation, Helmet.js headers, and API rate limiters.",
          "Enabled multilingual patient services in English, Hindi, and Gujarati, integrating Browser Web Speech APIs.",
          "Delivered professional developer workspaces contextually tracking appointments, schedules, and leaves."],
         COLOR_TEAL_600)

# Right Column - Thank you card (Solid colored callout slide)
thank_you_bg = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
thank_you_bg.fill.solid()
thank_you_bg.fill.fore_color.rgb = COLOR_TEAL_600
thank_you_bg.line.fill.background()

# Title "Thank You!"
tb_ty = slide10.shapes.add_textbox(Inches(7.2), Inches(2.6), Inches(4.9), Inches(1.0))
tf_ty = tb_ty.text_frame
tf_ty.word_wrap = True
p_ty = tf_ty.paragraphs[0]
p_ty.text = "Thank You"
p_ty.alignment = PP_ALIGN.CENTER
p_ty.font.name = "Calibri"
p_ty.font.size = Pt(44)
p_ty.font.bold = True
p_ty.font.color.rgb = COLOR_WHITE

# Subtitle
p_ty_sub = tf_ty.add_paragraph()
p_ty_sub.text = "Questions & Discussion"
p_ty_sub.alignment = PP_ALIGN.CENTER
p_ty_sub.font.name = "Calibri"
p_ty_sub.font.size = Pt(18)
p_ty_sub.font.color.rgb = COLOR_WHITE
p_ty_sub.space_before = Pt(8)

# Details
p_ty_det = tf_ty.add_paragraph()
p_ty_det.text = "\nMajor Project Presentation\nHealthAI Platform — 2026"
p_ty_det.alignment = PP_ALIGN.CENTER
p_ty_det.font.name = "Calibri"
p_ty_det.font.size = Pt(13)
p_ty_det.font.color.rgb = COLOR_SLATE_50
p_ty_det.space_before = Pt(20)

# Save the presentation
output_path = "HealthAI_Project_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")
